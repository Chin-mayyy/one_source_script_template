import os
import requests
from docx import Document
from pathlib import Path
import time
import re

class DocxUploader:
    def __init__(self, api_endpoint, auth_token, folder_path):
        self.api_endpoint = api_endpoint
        self.headers = {
            'Content-Type': 'application/json',
            'Authorization': f'Bearer {auth_token}'  # Adjust format as needed
        }
        self.folder_path = Path(folder_path)

    def extract_docx_content(self, file_path):
        """Extract text content from a .docx file"""
        try:
            doc = Document(file_path)
            content = []

            for paragraph in doc.paragraphs:
                if paragraph.text.strip():  # Skip empty paragraphs
                    content.append(paragraph.text.strip())
            return '\n'.join(content)

        except Exception as e:
            print(f"Error reading {file_path}: {str(e)}")
            return None

    def extract_bracket_content(self, text):
        """
        Extract all content inside {{}} and <<>> brackets from a string.

        Args:
            text (str): Input string containing bracketed content

        Returns:
            list: Array of strings containing the extracted content (without brackets)
        """
        text = str(text)
        # Extract content inside {{}}
        curly_braces = re.findall(r'\{\{(.*?)\}\}', text, re.DOTALL)

        # Extract content inside <<>>
        angle_brackets = re.findall(r'<<(.*?)>>', text, re.DOTALL)

        # Combine both lists and remove duplicates while preserving order
        all_content = []
        seen = set()

        for item in curly_braces + angle_brackets:
            # Strip whitespace and newlines
            cleaned_item = item.strip()
            if cleaned_item and cleaned_item not in seen:
                all_content.append(cleaned_item)
                seen.add(cleaned_item)

        return all_content

    def create_request_body(self, filename, content):
        """Create the request body structure"""
        # Extract name from filename (remove .docx extension)
        filename = Path(filename)
        name = filename.stem
        print(f"Processing file: {name}")

        # You can customize these fields based on your requirements
        markers = self.extract_bracket_content(content)
        request_body = {
            "formId": "68c99de867cffddcbec94f02",
            "name": name,
            "subject": f"{name}",
            "subjectPlaceholders": [],#will be empty for document templates
            "content": content,
            "placeholders": markers,
            "category": "",  # You can customize this
            "type": "document"
        }

        return request_body

    def upload_document(self, request_body, filename):
        """Upload a single document to the API"""
        try:
            response = requests.post(
                self.api_endpoint,
                headers=self.headers,
                json=request_body,
                timeout=30
            )

            if response.status_code == 200 or response.status_code == 201:
                print(f"✅ Successfully uploaded: {filename}")
                return True
            else:
                print(f"❌ Failed to upload {filename}. Status: {response.status_code}")
                print(f"Response: {response.text}")
                return False

        except requests.exceptions.RequestException as e:
            print(f"❌ Network error uploading {filename}: {str(e)}")
            return False

    def process_all_documents(self):
        """Process all .docx files in the folder"""
        # Find all .docx files
        docx_files = list(self.folder_path.glob("*.docx"))

        if not docx_files:
            print(f"No .docx files found in {self.folder_path}")
            return

        print(f"Found {len(docx_files)} .docx files to process")

        successful_uploads = 0
        failed_uploads = 0

        for file_path in docx_files:
            print(f"\nProcessing: {file_path.name}")

            # Skip temporary files (those starting with ~$)
            if file_path.name.startswith('~$'):
                print(f"Skipping temporary file: {file_path.name}")
                continue

            # Extract content from docx
            content = self.extract_docx_content(file_path)

            if content is None:
                failed_uploads += 1
                continue

            if not content.strip():
                print(f"⚠️  Warning: {file_path.name} appears to be empty")
                failed_uploads += 1
                continue

            # Create request body
            request_body = self.create_request_body(file_path, content)  # Pass the full file_path, not just the name

            # Upload document
            if self.upload_document(request_body, file_path.name):
                successful_uploads += 1
            else:
                failed_uploads += 1

            # Add small delay to avoid overwhelming the API
            time.sleep(0.5)

        # Summary
        print(f"\n{'='*50}")
        print(f"✅ Successful uploads: {successful_uploads}")
        print(f"❌ Failed uploads: {failed_uploads}")
        print(f"📁 Total files processed: {successful_uploads + failed_uploads}")


def main():
    # Configuration - UPDATE THESE VALUES
    API_ENDPOINT = "https://onesource.viithiisyserp.com/api/templates"  # Replace with your actual endpoint
    AUTH_TOKEN = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJhdXRoVXNlcklkIjoiNGY3MGEzODUtMGE5ZS00MTMzLTg3OTUtM2Q3NTFlMTA4Yzg4Iiwicm9sZUlkIjoiNGY2NmRiM2QtMjQyZC00NjEzLWE1M2UtMzk2MjY1YzJjYTRlIiwicm9sZU5hbWUiOiJBZG1pbmlzdHJhdG9yIiwiZW1haWwiOiJzaGFudHkuc2FpbmlAdmlpdGhpaXN5cy5jb20iLCJuYW1lIjoiU2hhbnR5IiwiaWF0IjoxNzU4ODIxNTA4LCJleHAiOjE3NTkwODA3MDh9.mbDzgXGfbheJCeZYnP-hSOyU83dRv5xSu6BkgdPKfL8"
    FOLDER_PATH = "moved_templates2/"  # Replace with your folder path

    # Validate inputs
    if AUTH_TOKEN == "your_auth_token_here":
        print("❌ Please update the AUTH_TOKEN in the script")
        return

    if API_ENDPOINT == "https://your-api-endpoint.com/upload":
        print("❌ Please update the API_ENDPOINT in the script")
        return

    if not os.path.exists(FOLDER_PATH):
        print(f"❌ Folder path does not exist: {FOLDER_PATH}")
        return

    # Create uploader instance and process documents
    uploader = DocxUploader(API_ENDPOINT, AUTH_TOKEN, FOLDER_PATH)
    uploader.process_all_documents()

if __name__ == "__main__":
    main()
